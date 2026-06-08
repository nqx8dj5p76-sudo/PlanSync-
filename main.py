from fastapi import FastAPI, UploadFile, File, Form, HTTPException
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import Response
import pdfplumber
import openpyxl
import re
import io
from datetime import datetime, timedelta

app = FastAPI(title="PlanSync API")

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_methods=["*"],
    allow_headers=["*"],
)

# ── Spalten-Grenzen (Mittelpunkte aus pdfplumber-Analyse) ─────────────────────
PDF_COLS = [
    ('früh',       60,  114),
    ('mitte',     114,  185),
    ('spät',      185,  253),
    ('nacht',     253,  317),
    ('1. dienst', 317,  391),
    ('2. dienst', 391,  473),
    ('prämed',    473,  543),
    ('op-spät',   543,  610),
    ('itw',       610,  700),
]

MO_MAP = {
    'jan':1,'feb':2,'mär':3,'mar':3,'apr':4,'mai':5,'jun':6,
    'jul':7,'aug':8,'sep':9,'okt':10,'nov':11,'dez':12
}

# ── Schichtzeiten ─────────────────────────────────────────────────────────────
def get_shift_times(col_key, date):
    dow = date.weekday()  # 0=Mo, 4=Fr, 5=Sa, 6=So
    is_weekend = dow >= 5
    is_friday  = dow == 4

    if col_key == 'früh':
        return ('Frühdienst', '07:00', '16:30', False, 'Schicht')
    elif col_key == 'mitte':
        return ('Spätdienst', '14:00', '22:30', False, 'Schicht')
    elif col_key == 'spät':
        return ('Spätdienst', '14:00', '22:30', False, 'Schicht')
    elif col_key == 'nacht':
        if is_weekend: return ('Nachtdienst', '20:00', '09:00', True,  'Schicht')
        if is_friday:  return ('Nachtdienst', '21:30', '09:00', True,  'Schicht')
        return             ('Nachtdienst', '21:30', '07:30', True,  'Schicht')
    elif col_key in ('1. dienst', '2. dienst'):
        label = '1. Dienst' if col_key == '1. dienst' else '2. Dienst'
        if is_weekend: return (label, '09:00', '09:00', True, 'Bereitschaft')
        return             (label, '09:30', '07:30', True, 'Bereitschaft')
    elif col_key == 'prämed':
        return ('Prämedikation', '16:00', '23:59', False, 'Sonderaufgabe')
    elif col_key == 'op-spät':
        return ('OP-Spät', '16:00', '18:00', False, 'Sonderaufgabe')
    elif col_key == 'itw':
        return ('ITW-Dienst', '07:00', '16:30', False, 'Schicht')
    return None

# ── PDF Parser ────────────────────────────────────────────────────────────────
def get_col(x):
    for name, xmin, xmax in PDF_COLS:
        if xmin <= x < xmax:
            return name
    return None

DAY_RE = re.compile(r'^(Mo|Di|Mi|Do|Fr|Sa|So)\.?$', re.I)

def parse_pdf(file_bytes, name_search):
    results = []
    name_lower = name_search.lower()

    with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
        page = pdf.pages[0]
        words = page.extract_words()

        # Monat + Jahr erkennen
        month, year = 0, datetime.now().year
        for w in words:
            ms = w['text'].lower()[:3]
            if ms in MO_MAP and len(w['text']) >= 3 and not w['text'][0].isdigit():
                month = MO_MAP[ms]
            if re.match(r'^\d{4}$', w['text']):
                year = int(w['text'])
        if not month:
            raise ValueError("Monat nicht erkannt. Bitte Dateiformat prüfen.")

        # Zeilen nach Y gruppieren (Toleranz 1px – pdfplumber ist exakt)
        lines = {}
        for w in words:
            y = round(w['top'])
            lines.setdefault(y, []).append(w)

        # Datenzeilen extrahieren
        rows = []
        for y, ws in sorted(lines.items()):
            ws.sort(key=lambda w: w['x0'])
            if len(ws) < 3: continue
            if not DAY_RE.match(ws[0]['text']): continue
            try:
                day = int(ws[1]['text'].rstrip('.'))
            except:
                continue
            if not (1 <= day <= 31): continue
            try:
                date = datetime(year, month, day)
            except:
                continue
            row = {'date': date}
            for w in ws[2:]:
                col = get_col(w['x0'])
                if col:
                    row[col] = row.get(col, '') + w['text']
            rows.append(row)

        # Name in Spalten suchen
        for row in rows:
            for col_name, _, _ in PDF_COLS:
                val = row.get(col_name, '').strip()
                if not val or val in ('---', '--', '-'): continue
                # OA-Duplikate überspringen (z.B. "Albeser/Albeser")
                parts = val.split('/')
                if len(parts) == 2 and parts[0].strip() == parts[1].strip():
                    continue
                for n in parts:
                    if name_lower in n.lower():
                        times = get_shift_times(col_name, row['date'])
                        if times:
                            title, start, end, overnight, source = times
                            end_date = row['date'] + timedelta(days=1) if overnight else row['date']
                            results.append({
                                'date': row['date'],
                                'end_date': end_date,
                                'title': title,
                                'start': start,
                                'end': end,
                                'source': source,
                            })
                        break

    return sorted(results, key=lambda x: x['date'])

# ── XLSX Parser ───────────────────────────────────────────────────────────────
XLSX_LEFT_COLS = list(range(2, 17, 2))  # 2,4,6,...,16 (0-indexed)

XLSX_BD_COLS = [
    (21, 'FD/TW',      'fdtw',   'Schicht'),
    (22, 'Kurzdienst', 'kurz',   'Schicht'),
    (24, 'Nacht-BD',   'bd',     'Bereitschaft'),
    (25, '1. Dienst',  'bd',     'Bereitschaft'),
    (26, '2. Dienst',  'bd',     'Bereitschaft'),
    (27, 'Prämedikation','praemed','Sonderaufgabe'),
    (28, 'OP-Spät',    'opspat', 'Sonderaufgabe'),
]

def get_bd_times(type_, date):
    dow = date.weekday()
    is_weekend = dow >= 5
    if type_ == 'fdtw':   return ('08:00', '21:00', False)
    if type_ == 'kurz':   return ('10:00', '16:00', False)
    if type_ == 'bd':
        if is_weekend:    return ('09:00', '09:00', True)
        return                   ('09:30', '07:30', True)
    if type_ == 'praemed': return ('16:00', '23:59', False)
    if type_ == 'opspat':  return ('16:00', '18:00', False)
    return ('09:30', '07:30', True)

def get_xlsx_shift_times(code, date):
    dow = date.weekday()
    is_weekend = dow >= 5
    is_friday  = dow == 4
    c = code.upper().strip()
    if c in ('FD', 'F', 'IT'):
        return ('Frühdienst' if c != 'IT' else 'IT-Dienst', '07:00', '16:30', False, 'Schicht')
    if c in ('SD', 'S'):
        return ('Spätdienst', '14:00', '22:30', False, 'Schicht')
    if c == 'ND':
        if is_weekend: return ('Nachtdienst', '20:00', '09:00', True, 'Schicht')
        if is_friday:  return ('Nachtdienst', '21:30', '09:00', True, 'Schicht')
        return             ('Nachtdienst', '21:30', '07:30', True, 'Schicht')
    if c == 'TW':
        return ('Tagwache', '08:00', '21:00', False, 'Schicht')
    if c == 'FD/TW':
        if is_weekend: return ('FD/TW', '08:00', '21:00', False, 'Schicht')
        return             ('FD/TW', '07:00', '16:30', False, 'Schicht')
    if c == 'K':
        return ('Kurzdienst', '10:00', '16:00', False, 'Schicht')
    if c in ('U',):
        return ('Urlaub', '00:00', '23:59', False, 'Schicht')
    return None

XLSX_SKIP = {'-', '---', 'X', 'x', ''}

def parse_xlsx(file_bytes, name_search):
    results = []
    name_lower = name_search.lower()
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)

    for ws in wb.worksheets:
        rows = list(ws.iter_rows(values_only=True))

        # Linker Bereich: Namensspalte finden (Zeilen 3-9)
        person_col = None
        for r in range(3, min(10, len(rows))):
            row = rows[r]
            for c in XLSX_LEFT_COLS:
                if c < len(row) and row[c] and name_lower in str(row[c]).lower():
                    person_col = c
                    break
            if person_col is not None:
                break

        # Linker Bereich: Schichten
        if person_col is not None:
            for row in rows:
                if not row or len(row) <= 1: continue
                date_val = row[1]
                if not isinstance(date_val, datetime): continue
                code = str(row[person_col]).strip() if person_col < len(row) and row[person_col] else ''
                if code in XLSX_SKIP: continue
                times = get_xlsx_shift_times(code, date_val)
                if not times: continue
                title, start, end, overnight, source = times
                end_date = date_val + timedelta(days=1) if overnight else date_val
                results.append({
                    'date': date_val,
                    'end_date': end_date,
                    'title': title,
                    'start': start,
                    'end': end,
                    'source': source,
                })

        # Rechter Bereich: Bereitschaftsdienste
        for row in rows:
            if not row or len(row) <= 20: continue
            date_val = row[20]
            if not isinstance(date_val, datetime): continue
            for col_idx, label, type_, source in XLSX_BD_COLS:
                if col_idx >= len(row) or not row[col_idx]: continue
                cell_val = str(row[col_idx]).strip()
                if len(cell_val) < 2: continue
                names = cell_val.split('/')
                for n in names:
                    if name_lower in n.lower():
                        start, end, overnight = get_bd_times(type_, date_val)
                        end_date = date_val + timedelta(days=1) if overnight else date_val
                        results.append({
                            'date': date_val,
                            'end_date': end_date,
                            'title': label,
                            'start': start,
                            'end': end,
                            'source': source,
                        })
                        break

    # Deduplizieren
    seen = set()
    unique = []
    for r in results:
        key = (r['date'].date(), r['title'])
        if key not in seen:
            seen.add(key)
            unique.append(r)

    return sorted(unique, key=lambda x: x['date'])

# ── ICS Generator ─────────────────────────────────────────────────────────────
def generate_ics(shifts):
    stamp = datetime.utcnow().strftime('%Y%m%dT%H%M%SZ')
    uid_base = int(datetime.utcnow().timestamp())
    lines = [
        'BEGIN:VCALENDAR',
        'VERSION:2.0',
        'PRODID:-//PlanSync//DE',
        'CALSCALE:GREGORIAN',
        'METHOD:PUBLISH',
        'X-WR-CALNAME:Dienstplan',
        'X-WR-TIMEZONE:Europe/Berlin',
    ]
    for i, s in enumerate(shifts):
        all_day  = s['start'] == '00:00' and s['end'] == '23:59'
        praemed  = s['end'] == '23:59' and s['start'] != '00:00'
        lines.append('BEGIN:VEVENT')
        lines.append(f"UID:ps-{uid_base}-{i}@plansync")
        lines.append(f"DTSTAMP:{stamp}")
        if all_day:
            lines.append(f"DTSTART;VALUE=DATE:{s['date'].strftime('%Y%m%d')}")
            lines.append(f"DTEND;VALUE=DATE:{s['date'].strftime('%Y%m%d')}")
        else:
            lines.append(f"DTSTART;TZID=Europe/Berlin:{s['date'].strftime('%Y%m%d')}T{s['start'].replace(':','')}00")
            end_d = s['date'] if praemed else s['end_date']
            end_t = '235900' if praemed else s['end'].replace(':', '') + '00'
            lines.append(f"DTEND;TZID=Europe/Berlin:{end_d.strftime('%Y%m%d')}T{end_t}")
        lines.append(f"SUMMARY:{s['title']}")
        lines.append(f"DESCRIPTION:{s['source']}")
        lines.append('END:VEVENT')
    lines.append('END:VCALENDAR')
    return '\r\n'.join(lines)

# ── Arbeitszeitverstöße ───────────────────────────────────────────────────────
def check_violations(shifts):
    violations = []
    sorted_shifts = sorted(shifts, key=lambda s: (s['date'], s['start']))
    for i, s in enumerate(sorted_shifts):
        # § 3 ArbZG: Max 10h (außer Bereitschaft)
        if s['source'] != 'Bereitschaft':
            h1 = int(s['start'].split(':')[0]) * 60 + int(s['start'].split(':')[1])
            h2 = int(s['end'].split(':')[0])   * 60 + int(s['end'].split(':')[1])
            if s['end_date'] > s['date']: h2 += 24 * 60
            duration = (h2 - h1) / 60
            if duration > 10:
                violations.append(f"§3 ArbZG: {s['date'].strftime('%d.%m')} {s['title']} – {duration:.1f}h (max. 10h)")
        # § 5 ArbZG: Mind. 11h Ruhezeit
        if i > 0:
            prev = sorted_shifts[i-1]
            prev_end = datetime.combine(prev['end_date'].date(), datetime.strptime(prev['end'], '%H:%M').time())
            cur_start = datetime.combine(s['date'].date(),       datetime.strptime(s['start'], '%H:%M').time())
            rest_h = (cur_start - prev_end).total_seconds() / 3600
            if 0 <= rest_h < 11:
                violations.append(f"§5 ArbZG: {prev['date'].strftime('%d.%m')} → {s['date'].strftime('%d.%m')} – nur {rest_h:.1f}h Ruhezeit (min. 11h)")
    return violations

# ── API Endpoints ─────────────────────────────────────────────────────────────
@app.get("/")
def root():
    return {"status": "PlanSync API läuft", "version": "1.0"}

@app.post("/extract")
async def extract(
    file: UploadFile = File(...),
    name: str = Form(...)
):
    if not name.strip():
        raise HTTPException(400, "Name darf nicht leer sein")

    file_bytes = await file.read()
    filename = file.filename.lower()

    try:
        if filename.endswith('.pdf'):
            shifts = parse_pdf(file_bytes, name)
        elif filename.endswith(('.xlsx', '.xls')):
            shifts = parse_xlsx(file_bytes, name)
        else:
            raise HTTPException(400, "Nur PDF oder XLSX Dateien erlaubt")
    except ValueError as e:
        raise HTTPException(422, str(e))
    except Exception as e:
        raise HTTPException(500, f"Fehler beim Lesen der Datei: {str(e)}")

    if not shifts:
        raise HTTPException(404, f"Keine Einträge für '{name}' gefunden")

    ics_content = generate_ics(shifts)
    violations  = check_violations(shifts)

    return Response(
        content=ics_content,
        media_type="text/calendar",
        headers={
            "Content-Disposition": "attachment; filename=PlanSync.ics",
            "X-Shifts-Count":      str(len(shifts)),
            "X-Violations":        "; ".join(violations) if violations else "",
            "Access-Control-Expose-Headers": "X-Shifts-Count, X-Violations",
        }
    )

@app.post("/names")
async def get_names(file: UploadFile = File(...)):
    """Gibt alle Namen aus der Datei zurück"""
    file_bytes = await file.read()
    filename = file.filename.lower()
    names = set()

    try:
        if filename.endswith('.pdf'):
            with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
                page = pdf.pages[0]
                words = page.extract_words()
                month, year = 0, datetime.now().year
                for w in words:
                    ms = w['text'].lower()[:3]
                    if ms in MO_MAP and len(w['text']) >= 3:
                        month = MO_MAP[ms]
                    if re.match(r'^\d{4}$', w['text']):
                        year = int(w['text'])
                if month:
                    lines = {}
                    for w in words:
                        y = round(w['top'])
                        lines.setdefault(y, []).append(w)
                    for y, ws in sorted(lines.items()):
                        ws.sort(key=lambda w: w['x0'])
                        if len(ws) < 3 or not DAY_RE.match(ws[0]['text']): continue
                        for w in ws[2:]:
                            col = get_col(w['x0'])
                            if col:
                                for n in w['text'].split('/'):
                                    t = n.strip()
                                    if len(t) >= 2 and not t[0].isdigit() and t not in ('---','--'):
                                        names.add(t)
        elif filename.endswith(('.xlsx', '.xls')):
            wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
            for ws in wb.worksheets:
                rows = list(ws.iter_rows(values_only=True))
                for r in range(3, min(10, len(rows))):
                    row = rows[r]
                    for c in XLSX_LEFT_COLS:
                        if c < len(row) and row[c]:
                            v = str(row[c]).strip()
                            if v and not v[0].isdigit() and 'KW' not in v:
                                names.add(v)
    except Exception as e:
        raise HTTPException(500, str(e))

    return {"names": sorted(names)}
